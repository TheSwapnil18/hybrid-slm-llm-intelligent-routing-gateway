# build_final_report.py - Master Compiler for the Hybrid SLM-LLM Router Agent Project Report
import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

# Ensure local Report directory is in python search path
REPORT_DIR = os.path.dirname(os.path.abspath(__file__))
if REPORT_DIR not in sys.path:
    sys.path.insert(0, REPORT_DIR)

from front_matter import build_front_matter
from chapter1 import build_chapter1
from chapter2 import build_chapter2
from chapter3 import build_chapter3
from chapter4 import build_chapter4
from chapter5 import (
    build_chapter5_content, build_references_content, build_acknowledgement_content
)

print("Starting generation of comprehensive report_sem7.docx...")

doc = docx.Document()

# -------------------------------------------------------------
# XML HELPER FUNCTIONS
# -------------------------------------------------------------
def set_cell_margins(cell, top=60, bottom=60, left=100, right=100):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        r'<w:tcMar {}><w:top w:w="{}" w:type="dxa"/><w:bottom w:w="{}" w:type="dxa"/>'
        r'<w:left w:w="{}" w:type="dxa"/><w:right w:w="{}" w:type="dxa"/></w:tcMar>'.format(
            nsdecls('w'), top, bottom, left, right
        )
    )
    tcPr.append(tcMar)

def add_footer_page_number(footer_para):
    footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_para.text = ""
    r = footer_para.add_run()
    fldSimple = OxmlElement('w:fldSimple')
    fldSimple.set(qn('w:instr'), 'PAGE')
    r._r.append(fldSimple)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor(0, 0, 0)

def setup_section_margins(section):
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.5)
    section.right_margin = Inches(0.5)

def setup_header(section, chapter_title):
    header = section.header
    header.is_linked_to_previous = False
    p_hdr = header.paragraphs[0]
    p_hdr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_hdr.text = ""
    r = p_hdr.add_run(chapter_title)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor(0, 0, 0)

def setup_footer(section, restart_numbering=False):
    footer = section.footer
    footer.is_linked_to_previous = False
    p_ft = footer.paragraphs[0]
    add_footer_page_number(p_ft)
    
    if restart_numbering:
        sectPr = section._sectPr
        pgNumType = OxmlElement('w:pgNumType')
        pgNumType.set(qn('w:start'), '1')
        pgNumType.set(qn('w:fmt'), 'decimal')
        sectPr.append(pgNumType)

# -------------------------------------------------------------
# CONFIGURE DEFAULT STYLES (Strict TNR 12 Justified Pure Black)
# -------------------------------------------------------------
style_normal = doc.styles['Normal']
font = style_normal.font
font.name = 'Times New Roman'
font.size = Pt(12)
font.color.rgb = RGBColor(0, 0, 0)
style_normal.paragraph_format.line_spacing = 1.15
style_normal.paragraph_format.space_after = Pt(6)
style_normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

# -------------------------------------------------------------
# TYPOGRAPHY BUILDER FUNCTIONS
# -------------------------------------------------------------
def add_p(text="", align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_before=0, space_after=6, bold=False, italic=False, size=12, color=None):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    if text:
        r = p.add_run(text)
        r.bold = bold
        r.italic = italic
        r.font.name = 'Times New Roman'
        r.font.size = Pt(size)
        r.font.color.rgb = color if color is not None else RGBColor(0, 0, 0)
    return p

def add_h1(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(14)
    r = p.add_run(text.upper())
    r.bold = True
    r.font.name = 'Times New Roman'
    r.font.size = Pt(16)
    r.font.color.rgb = RGBColor(0, 0, 0)
    return p

def add_h2(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    r.bold = True
    r.font.name = 'Times New Roman'
    r.font.size = Pt(14)
    r.font.color.rgb = RGBColor(0, 0, 0)
    return p

def add_h3(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    r.bold = True
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor(0, 0, 0)
    return p

def add_h4(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text)
    r.bold = True
    r.italic = True
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor(0, 0, 0)
    return p

def add_image_centered(img_path, caption, width_in=6.0):
    resolved_path = img_path
    if not os.path.exists(resolved_path):
        resolved_path = os.path.join(REPORT_DIR, os.path.basename(img_path))
    if not os.path.exists(resolved_path):
        resolved_path = os.path.join(REPORT_DIR, img_path)

    if os.path.exists(resolved_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(10)
        p_img.paragraph_format.space_after = Pt(4)
        run_img = p_img.add_run()
        run_img.add_picture(resolved_path, width=Inches(width_in))
        
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(2)
        p_cap.paragraph_format.space_after = Pt(12)
        r_cap = p_cap.add_run(caption)
        r_cap.bold = True
        r_cap.italic = True
        r_cap.font.name = 'Times New Roman'
        r_cap.font.size = Pt(10.5)
        r_cap.font.color.rgb = RGBColor(0, 0, 0)
    else:
        print(f"Warning: image path not found: {img_path} (resolved: {resolved_path})")

# -------------------------------------------------------------
# SIMPLE TABLE FORMATTER (NO DESIGN, NO BORDERS, PURE BLACK)
# -------------------------------------------------------------
def format_simple_table(table, col_widths=None):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    # Remove all borders from table
    tblPr = table._tbl.tblPr
    tblBorders = parse_xml(
        r'<w:tblBorders {}><w:top w:val="none"/><w:left w:val="none"/><w:bottom w:val="none"/>'
        r'<w:right w:val="none"/><w:insideH w:val="none"/><w:insideV w:val="none"/></w:tblBorders>'.format(nsdecls('w'))
    )
    tblPr.append(tblBorders)
    
    # Iterate through all cells: ensure no background shading, text is black Times New Roman
    for r_idx, row in enumerate(table.rows):
        for cell in row.cells:
            # Remove any background shading
            tcPr = cell._tc.get_or_add_tcPr()
            shd = tcPr.find(qn('w:shd'))
            if shd is not None:
                tcPr.remove(shd)
            set_cell_margins(cell, top=60, bottom=60, left=100, right=100)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(1)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.15
                for r in p.runs:
                    r.font.name = 'Times New Roman'
                    r.font.color.rgb = RGBColor(0, 0, 0)
                    if r_idx == 0:
                        r.bold = True
                        r.font.size = Pt(11)
                    else:
                        r.font.size = Pt(10.5)

    # Apply column widths
    if col_widths:
        for row in table.rows:
            for i, w in enumerate(col_widths):
                if i < len(row.cells):
                    row.cells[i].width = Inches(w)

# =============================================================
# COMPILE DOCUMENT SECTIONS
# =============================================================

# 1. SECTION 1: FRONT MATTER (Preliminaries)
print("1/7 Compiling Front Matter...")
s1 = doc.sections[0]
setup_section_margins(s1)
s1.different_first_page_header_footer = True

# Lowercase Roman numbering for front matter
sectPr1 = s1._sectPr
pgNumType1 = OxmlElement('w:pgNumType')
pgNumType1.set(qn('w:fmt'), 'lowerRoman')
sectPr1.append(pgNumType1)

add_footer_page_number(s1.footer.paragraphs[0])

build_front_matter(doc, add_p, add_h1, add_image_centered, format_simple_table)

# 2. SECTION 2: CHAPTER 1: INTRODUCTION
print("2/7 Compiling Chapter 1: Introduction...")
s_ch1 = doc.add_section(WD_SECTION_START.CONTINUOUS)
setup_section_margins(s_ch1)
s_ch1.different_first_page_header_footer = True
setup_header(s_ch1, "Chapter 1: Introduction")
setup_footer(s_ch1, restart_numbering=True)

build_chapter1(doc, add_p, add_h1, add_h2, add_h3, add_image_centered)

# 3. SECTION 3: CHAPTER 2: LITERATURE SURVEY
print("3/7 Compiling Chapter 2: Literature Survey...")
s_ch2 = doc.add_section(WD_SECTION_START.CONTINUOUS)
setup_section_margins(s_ch2)
setup_header(s_ch2, "Chapter 2: Literature Survey")
setup_footer(s_ch2, restart_numbering=False)

build_chapter2(doc, add_p, add_h1, add_h2, add_h3, format_simple_table)

# 4. SECTION 4: CHAPTER 3: SYSTEM DEVELOPMENT
print("4/7 Compiling Chapter 3: System Development...")
s_ch3 = doc.add_section(WD_SECTION_START.CONTINUOUS)
setup_section_margins(s_ch3)
setup_header(s_ch3, "Chapter 3: System Development")
setup_footer(s_ch3, restart_numbering=False)

build_chapter3(doc, add_p, add_h1, add_h2, add_h3, add_h4, add_image_centered, format_simple_table)

# 5. SECTION 5: CHAPTER 4: PERFORMANCE EVALUATION
print("5/7 Compiling Chapter 4: Performance Evaluation...")
s_ch4 = doc.add_section(WD_SECTION_START.CONTINUOUS)
setup_section_margins(s_ch4)
setup_header(s_ch4, "Chapter 4: Performance Evaluation")
setup_footer(s_ch4, restart_numbering=False)

build_chapter4(doc, add_p, add_h1, add_h2, add_h3, add_h4, add_image_centered, format_simple_table)

# 6. SECTION 6: CHAPTER 5: CONCLUSION & FUTURE SCOPE
print("6/7 Compiling Chapter 5: Conclusion & Future Scope...")
s_ch5 = doc.add_section(WD_SECTION_START.CONTINUOUS)
setup_section_margins(s_ch5)
setup_header(s_ch5, "Chapter 5: Conclusion & Future Scope")
setup_footer(s_ch5, restart_numbering=False)

build_chapter5_content(doc, add_p, add_h1, add_h2, add_h3)

# 7. SECTION 7: REFERENCES & ACKNOWLEDGEMENTS
print("7/7 Compiling References & Acknowledgements...")
s_ref = doc.add_section(WD_SECTION_START.CONTINUOUS)
setup_section_margins(s_ref)
setup_header(s_ref, "References")
setup_footer(s_ref, restart_numbering=False)

build_references_content(doc, add_p, add_h1)

s_ack = doc.add_section(WD_SECTION_START.CONTINUOUS)
setup_section_margins(s_ack)
setup_header(s_ack, "Acknowledgement")
setup_footer(s_ack, restart_numbering=False)

build_acknowledgement_content(doc, add_p, add_h1)

# -------------------------------------------------------------
# SAVE FINAL REPORT
# -------------------------------------------------------------
target_file = os.path.join(REPORT_DIR, "report_sem7.docx")
print(f"Saving compiled master report to: {target_file}")
doc.save(target_file)
print(f"Master report successfully saved! File size: {os.path.getsize(target_file):,} bytes")
